import importlib.util
from pathlib import Path
from openpyxl import load_workbook
spec=importlib.util.spec_from_file_location('campaign_summary_candidate',Path(__file__).parents[1]/'campaign_result_summary.py')
m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)


def test_batch_success_not_full_sku_success(tmp_path):
    r={'status':'complete','all_signed_up':False,'segments':[{'segment_id':'s','success':{'123':'receipt'},'exceptions':{'456':[{'reason':'缺少价格基线'}]}}],
       'final_audit':{'segments':[{'segment_id':'s','products':[{'item':'123','status':'sku_scope_unconfirmed'},{'item':'456','status':'not_verified_registered'}]}]}}
    rows=m.rows_for(r);assert rows[0]['execution']=='成功记录已确认' and rows[0]['official_readback']=='存在未确认SKU范围'
    a=m.write_summary(r,tmp_path);assert m.write_summary(r,tmp_path)==a
    book=load_workbook(a['xlsx_path']);assert book['报名结果']['A2'].value=='123'
    assert book['报名结果']['D3'].value=='缺少价格基线'


def test_blocked_without_audit_not_reported_success():
    rows=m.rows_for({'status':'blocked','segments':[{'exceptions':{'1':[{'reason':'unknown'}]}}]})
    assert rows[0]['official_readback']=='未取得完整回读' and rows[0]['execution']=='仍有待处理项'


def test_external_text_is_not_excel_formula(tmp_path):
    r={'status':'complete','segments':[{'exceptions':{'1':[{'reason':'=1+1'}]}}]}
    result=m.write_summary(r,tmp_path);book=load_workbook(result['xlsx_path'])
    assert book['报名结果']['D2'].value=='=1+1' and book['报名结果']['D2'].data_type=='s'
