"""One local end-of-run report, preserving batch and current SKU states separately."""
import hashlib
import json
from pathlib import Path

REASONS={
    'erp_mapping_missing_or_not_unique':'淘宝SKU缺少唯一ERP对应关系',
    'mapping_missing_or_ambiguous':'商品或SKU对应关系不明确',
    'fixed_original_evidence_missing':'缺少定制SKU首次原价基线，不能自动降价',
    'fixed_basis_unknown_not_rotation_proof':'缺少定制SKU原价基线，不能推定需要轮换',
    'unparsed_or_incomplete_official_failure':'官方失败说明未能完整解析，需补齐事实',
    'below_fixed_twenty_percent':'所需价格低于首次原价20%底线，需决定是否轮换',
    'final_delta_over_two_yuan':'超过同一目标价允许的2元修正范围，需决定是否轮换',
    'price_evidence_incomplete':'缺少完整价格事实',
    'verified_price_adjustment_had_no_effect_on_official_price':'已修正优惠但平台反馈价格未变化，不能重复盲改',
    'platform_requires_non_daily_signup_price':'平台要求与ERP日常报名价不一致，不能自行改口径',
}


def rows_for(result):
    audits={s.get('segment_id'):s for s in (result.get('final_audit') or {}).get('segments',[])}
    rows=[]
    labels={'registered':'完整核对范围已报名','sku_scope_unconfirmed':'存在未确认SKU范围',
            'not_verified_registered':'尚未确认报名','not_registered':'未报名'}
    for segment in result.get('segments',[]):
        audit=audits.get(segment.get('segment_id'),{})
        products={p['item']:p for p in audit.get('products',[])}
        success=segment.get('success',{});exceptions=segment.get('exceptions',{})
        for item in sorted(set(success)|set(exceptions)|set(products)):
            issues=exceptions.get(item,[]);p=products.get(item,{})
            rows.append({'item':str(item),'segment_id':segment.get('segment_id'),
                'execution':'成功记录已确认' if item in success else '仍有待处理项' if issues else '仅有回读记录',
                'official_readback':labels.get(p.get('status'),'未取得完整回读'),
                'issues':'；'.join(dict.fromkeys(REASONS.get(str(e.get('reason')),str(e.get('reason') or e.get('action') or '待核对')) for e in issues)),
                'issue_codes':[e.get('reason') or e.get('action') for e in issues]})
    return rows


def write_summary(result,root):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    rows=rows_for(result)
    summary={'schema':'campaign-two-upload-summary-v1','controller_status':result.get('status'),
        'all_signed_up':result.get('all_signed_up') is True,'rows':rows,
        'note':'成功记录与当前完整SKU核对分列；不能将批次成功、部分SKU成功或本地测试通过等同全店全部报名成功。'}
    raw=json.dumps(summary,ensure_ascii=False,sort_keys=True).encode('utf-8')
    digest=hashlib.sha256(raw).hexdigest();folder=Path(root)/'summary'/digest;folder.mkdir(parents=True,exist_ok=True)
    jp=folder/'报名结果汇总.json';xp=folder/'报名结果汇总.xlsx'
    if jp.exists() and jp.read_bytes()!=raw:raise ValueError('summary_content_changed')
    if not jp.exists():jp.write_bytes(raw)
    if not xp.exists():
        book=Workbook();sheet=book.active;sheet.title='报名结果'
        sheet.append(['商品ID','执行记录','平台当前回读','待处理原因'])
        for row in rows:sheet.append([row['item'],row['execution'],row['official_readback'],row['issues']])
        for cell in sheet[1]:
            cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='24476B')
        for line in sheet.iter_rows(min_row=2):
            line[0].number_format='@'
            for cell in line:cell.alignment=Alignment(vertical='top',wrap_text=True)
        for col,width in [('A',20),('B',22),('C',28),('D',85)]:sheet.column_dimensions[col].width=width
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        notes=book.create_sheet('口径');notes.append(['说明',summary['note']]);notes.column_dimensions['B'].width=110
        notes['B1'].alignment=Alignment(wrap_text=True)
        book.save(xp)
    return {'json_path':str(jp),'json_sha256':digest,'xlsx_path':str(xp),
            'xlsx_sha256':hashlib.sha256(xp.read_bytes()).hexdigest(),'platform_write':False}
