"""Current ERP listing state gate. Read only; never edits upload bytes or prices."""
from __future__ import annotations

import hashlib
import json
import re
import zipfile
from xml.etree import ElementTree
from datetime import datetime, timezone
from io import BytesIO

import openpyxl
from sqlalchemy import select

from app.models.product import Product
from app.models.pricing import PricingSku
from app.models.pricing_ext import PricingSkuPromo
from app.models.taobao_listing import TaobaoListing
from app.models.settings import SystemSetting

CHANNELS = frozenset({'product_prices', 'product_info', 'product_upload',
                      'single_item_discount', 'promo_signup', 'super_reduce'})
ACTIVE = frozenset({'上架', '在售', '出售中'})


def text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def norm(value):
    return re.sub(r'[\s_＊*：:（）()\-]', '', text(value)).lower()


class ListingGuardError(ValueError):
    pass


def read_registry(db):
    raw = db.execute(select(SystemSetting.value_plain).where(
        SystemSetting.key == 'product_sheet_sku_status_overrides_v1')).scalar_one_or_none()
    # Missing is an unavailable authority, not an empty exclusion list.
    if not raw:
        raise ListingGuardError('sku_listing_registry_unavailable')
    try:
        data = json.loads(raw)
        records = data['records']
        if not isinstance(records, list):
            raise ValueError()
        for row in records:
            if (not isinstance(row, dict) or not row.get('status') or
                    not (row.get('erp_sku_code') or
                         (row.get('product_code') and row.get('source_sku')))):
                raise ValueError()
        return records
    except (ValueError, KeyError, TypeError) as exc:
        raise ListingGuardError('sku_listing_registry_invalid') from exc


def sku_block_reason(sku, product_status, records):
    if text(product_status) not in ACTIVE:
        return 'product_offsale_or_status_unknown'
    for row in records:
        same = (row.get('erp_sku_code') == sku.sku_code or
                (row.get('product_code') == sku.product_code and
                 norm(row.get('source_sku')) == norm(sku.sku)))
        if same and row['status'] not in ACTIVE:
            return 'sku_offsale'
    if 'ERP规格状态：下架' in (sku.remark or ''):
        return 'sku_offsale'
    return ''


def snapshot(db):
    """Column queries bypass ORM identity-map caches, including old prepared jobs."""
    from types import SimpleNamespace
    records = read_registry(db)
    skus = {r.sku_code: SimpleNamespace(**r._asdict()) for r in db.execute(select(
        PricingSku.sku_code, PricingSku.product_code, PricingSku.sku, PricingSku.remark))}
    statuses = dict(db.execute(select(Product.code, Product.listing_status)).all())
    pairs = {}
    for r in db.execute(select(PricingSkuPromo.sku_code, PricingSkuPromo.taobao_item_id,
                               PricingSkuPromo.taobao_sku_id, PricingSkuPromo.alt_taobao_sku_ids)):
        for sid in [r.taobao_sku_id, *(r.alt_taobao_sku_ids or [])]:
            if text(r.taobao_item_id).isdigit() and text(sid).isdigit():
                pairs.setdefault((text(r.taobao_item_id), text(sid)), set()).add(r.sku_code)
    current_pairs = set(pairs)
    for r in db.execute(select(TaobaoListing.taobao_item_id, TaobaoListing.taobao_sku_id,
                               TaobaoListing.sku_code)):
        if (r.sku_code and text(r.taobao_sku_id).isdigit() and
                (text(r.taobao_item_id), text(r.taobao_sku_id)) not in current_pairs):
            # The imported listing is a historical snapshot. An explicit current
            # ERP mapping wins over its old merchant code after a reconciliation.
            pairs.setdefault((text(r.taobao_item_id), text(r.taobao_sku_id)), set()).add(r.sku_code)
    raw = db.execute(select(SystemSetting.value_plain).where(
        SystemSetting.key == 'delisted_skuids')).scalar_one_or_none() or '[]'
    try:
        ids = json.loads(raw)
        if not isinstance(ids, list):
            raise ValueError()
        excluded = {text(x) for x in ids}
    except (TypeError, ValueError) as exc:
        raise ListingGuardError('delisted_sku_registry_invalid') from exc
    for r in records:
        if r['status'] not in ACTIVE:
            excluded.update(map(text, r.get('taobao_sku_ids', [])))
    return skus, statuses, pairs, excluded, records


def validate_rows(db, rows):
    skus, statuses, pairs, excluded, records = snapshot(db)
    blocked, checked = [], []
    for row in rows:
        item, sid, code = (text(row.get(k)) for k in ('item_id', 'sku_id', 'sku_code'))
        candidates = pairs.get((item, sid), set()) if sid else set()
        reason = ''
        if sid in excluded:
            reason = 'taobao_sku_offsale'
        elif sid and (len(candidates) != 1 or (code and code not in candidates)):
            reason = 'sku_identity_unresolved_or_conflicting'
        elif not sid and code in skus:
            # New products without platform IDs still have an exact ERP SKU key.
            candidates = {code}
            if item and not any(i == item and code in codes for (i, _), codes in pairs.items()):
                reason = 'item_identity_unresolved_or_conflicting'
        elif not sid:
            reason = 'erp_sku_code_required'
        if not reason:
            code = next(iter(candidates))
            sku = skus.get(code)
            reason = (sku_block_reason(sku, statuses.get(sku.product_code), records)
                      if sku else 'erp_sku_missing')
        fact = dict(row, sku_code=code, reason=reason)
        (blocked if reason else checked).append(fact)
    return {'ok': bool(rows) and not blocked, 'error': 'upload_listing_guard_blocked' if blocked else
            '' if rows else 'upload_scope_empty', 'checked_count': len(checked), 'blocked': blocked,
            'checked_at': datetime.now(timezone.utc).isoformat(), 'checked': checked}


def workbook_rows(data):
    """Read official data sheet/header, never modify ZIP/custom XML or prices."""
    try:
        wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        raise ListingGuardError('upload_workbook_unreadable') from exc
    found = []
    try:
        for ws in wb.worksheets:
            header = None
            header_number = 0
            merge_ranges = []
            anchors = {}
            with zipfile.ZipFile(BytesIO(data)) as archive:
                for _, element in ElementTree.iterparse(archive.open(ws._worksheet_path)):
                    if element.tag.endswith('}mergeCell'):
                        merge_ranges.append(openpyxl.utils.range_boundaries(element.attrib['ref']))
                    element.clear()
            for number, values in enumerate(ws.iter_rows(values_only=True), 1):
                labels = [norm(re.split(r'[（(]', text(v).split('\n')[0])[0]) for v in values]
                if header is None:
                    if number > 30:
                        break
                    item_cols = [i for i, v in enumerate(labels) if v in
                                 {'商品id', '宝贝id', 'itemid', '商品数字id'}]
                    sku_cols = [i for i, v in enumerate(labels) if v in {'skuid', 'skuid必填'}]
                    code_cols = [i for i, v in enumerate(labels) if v in
                                 {'商家编码', 'sku商家编码', 'skucode', 'erp商家编码'}]
                    if (sku_cols or code_cols) and item_cols:
                        header = (item_cols[0], sku_cols[0] if sku_cols else None,
                                  code_cols[-1] if code_cols else None)
                        header_number = number
                    continue
                if not any(v is not None and text(v) for v in values):
                    continue
                item_col, sku_col, code_col = header
                def get(i):
                    return text(values[i]) if i is not None and i < len(values) else ''
                item, sid, code = get(item_col), get(sku_col), get(code_col)
                if item:
                    anchors[number] = item
                else:
                    for c1, r1, c2, r2 in merge_ranges:
                        if c1 == c2 == item_col + 1 and r1 < number <= r2:
                            item = anchors.get(r1, '')
                            break
                if (number == header_number + 1 and ws.title == '商品SKU导入列表'
                        and item.startswith('必填') and '输入报名商品的SKUID' in sid):
                    continue  # Official instruction row, not a data SKU.
                if any(v.startswith('=') for v in (item, sid, code)):
                    raise ListingGuardError('formula_in_upload_identity')
                # A nonempty price row with missing identity must not disappear
                # behind valid siblings. Blank rows were already skipped above.
                found.append({'sheet': ws.title, 'row': number, 'item_id': item,
                              'sku_id': sid, 'sku_code': code})
    finally:
        wb.close()
    if not found:
        raise ListingGuardError('upload_header_or_rows_unrecognized')
    return found


def validate_workbook(db, data):
    result = validate_rows(db, workbook_rows(data))
    result['sha256'] = hashlib.sha256(data).hexdigest()
    return result
