import json
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock

import openpyxl
import pytest
from sqlalchemy import update

from app.models.product import Product
from app.models.pricing import PricingSku
from app.models.pricing_ext import PricingSkuPromo
from app.models.settings import SystemSetting
from app.services import taobao_listing_guard_service as guard, web_agent_service as agent


@pytest.fixture
def ready(db_session):
    db = db_session
    db.add(Product(code='P', name='同款圆桌', listing_status='在售'))
    db.add_all([PricingSku(sku_code='P11', product_code='P', sku='95厘米'),
                PricingSku(sku_code='P12', product_code='P', sku='120厘米'),
                PricingSku(sku_code='P13', product_code='P', sku='未来规格')])
    db.add_all([PricingSkuPromo(sku_code='P11', taobao_item_id='123456', taobao_sku_id='111111',
                               alt_taobao_sku_ids=['333333']),
                PricingSkuPromo(sku_code='P12', taobao_item_id='123456', taobao_sku_id='222222')])
    db.add(SystemSetting(key='product_sheet_sku_status_overrides_v1', value_plain=json.dumps({'records':[
        {'erp_sku_code':'P11','status':'下架'},
        {'product_code':'P','source_sku':'未来规格','status':'下架'}]})))
    db.commit()
    return db


def workbook(rows, headers=None):
    wb=openpyxl.Workbook();ws=wb.active
    ws.append(headers or ['商品id','SKU_ID','优惠值'])
    for row in rows: ws.append(row)
    buf=BytesIO();wb.save(buf);return buf.getvalue()


def test_offsale_exact_scope_and_normal_sibling(ready):
    result=guard.validate_workbook(ready,workbook([['123456','111111',10],['123456','222222',20]]))
    assert result['ok'] is False
    assert result['checked_count']==1
    assert result['checked'][0]['sku_code']=='P12'
    assert result['blocked'][0]['sku_code']=='P11'


@pytest.mark.parametrize('sku', ['111111','333333'])
def test_primary_and_rotation_alias_both_blocked(ready,sku):
    assert not guard.validate_workbook(ready,workbook([['123456',sku,10]]))['ok']


def test_new_product_without_taobao_id_and_future_row(ready):
    headers=['商品id','商家编码','skuId','商家编码','价格(元)']
    assert guard.validate_workbook(ready,workbook([['','P','','P12',999]],headers))['ok']
    assert not guard.validate_workbook(ready,workbook([['','P','','P13',999]],headers))['ok']
    assert not guard.validate_workbook(ready,workbook([['','P','','',999]],headers))['ok']


def test_changed_state_rechecks_old_file_and_cached_product(ready):
    data=workbook([['123456','222222',20]])
    ready.query(Product).filter_by(code='P').one()
    assert guard.validate_workbook(ready,data)['ok']
    ready.execute(update(Product).values(listing_status='下架').execution_options(synchronize_session=False))
    ready.commit()
    assert not guard.validate_workbook(ready,data)['ok']


@pytest.mark.parametrize('rows',[[['654321','222222',10]],[['123456','999999',10]]])
def test_wrong_item_or_unknown_sku_never_guessed(ready,rows):
    assert not guard.validate_workbook(ready,workbook(rows))['ok']


def test_conflicting_merchant_code_blocked(ready):
    data=workbook([['123456','222222','P11']],['商品id','SKU_ID','商家编码'])
    assert not guard.validate_workbook(ready,data)['ok']
    data=workbook([['654321','','P12']],['商品id','SKU_ID','商家编码'])
    assert not guard.validate_workbook(ready,data)['ok']


@pytest.mark.parametrize('data', [b'not-xlsx',workbook([]),workbook([['123456','=1+1',10]])])
def test_malformed_or_formula_scope_blocked(ready,data):
    with pytest.raises(guard.ListingGuardError): guard.validate_workbook(ready,data)


@pytest.mark.parametrize('channel',['product_prices','single_item_discount','promo_signup','super_reduce'])
@pytest.mark.parametrize('phase',['stage','commit'])
def test_every_public_upload_stops_before_network(ready,monkeypatch,channel,phase):
    network=Mock(side_effect=AssertionError('must not upload'))
    monkeypatch.setattr(agent,'ensure_online',network)
    monkeypatch.setattr(agent.requests,'post',network)
    result=agent.upload_file(ready,channel,phase,workbook([['123456','111111',10]]),'old.xlsx')
    assert not result['ok'];network.assert_not_called()


@pytest.mark.parametrize('channel',['product_prices','single_item_discount','promo_signup','super_reduce'])
def test_normal_sku_uploaded_byte_identical(ready,monkeypatch,channel):
    data=workbook([['123456','222222',20]])
    monkeypatch.setattr(agent,'ensure_online',lambda *a,**k:{'online':True})
    response=Mock(status_code=200);response.json.return_value={'ok':True,'job':'test-only'}
    network=Mock(return_value=response);monkeypatch.setattr(agent.requests,'post',network)
    assert agent.upload_file(ready,channel,'commit',data,'ok.xlsx')['ok']
    assert network.call_args.kwargs['files']['file'][1]==data


def test_browser_wake_state_change_blocks_before_post(ready,monkeypatch):
    def wake(*a,**kw):
        ready.execute(update(Product).values(listing_status='下架'));ready.commit()
        return {'online':True}
    monkeypatch.setattr(agent,'ensure_online',wake)
    network=Mock();monkeypatch.setattr(agent.requests,'post',network)
    assert not agent.upload_file(ready,'promo_signup','commit',workbook([['123456','222222',20]]),'old.xlsx')['ok']
    network.assert_not_called()


def test_bad_registry_fails_closed(ready):
    ready.execute(update(SystemSetting).values(value_plain='broken'));ready.commit()
    with pytest.raises(guard.ListingGuardError): guard.snapshot(ready)


def test_generators_exclude_only_offsale(ready):
    from app.services.taobao_export_service import _fetch_rows
    from app.services.delisted_sku_service import get_delisted
    assert [s.sku_code for s,_,_ in _fetch_rows(ready,None,None)]==['P12']
    assert get_delisted(ready)=={'111111','333333'}


def test_official_multiline_discount_header(ready):
    data=workbook([['123456','222222',10]],['商品id\n(请使用文本格式)',
        'SKU_ID\n(活动为SKU级别必填，商品级别不填,一行只能填写一个sku)','优惠值：\n填写金额'])
    assert guard.validate_workbook(ready,data)['ok']


def test_official_signup_instruction_row(ready):
    wb=openpyxl.Workbook();ws=wb.active;ws.title='商品SKU导入列表'
    ws.append(['基础信息',None,None]);ws.append(['商品ID','SKUID','活动价'])
    ws.append(['必填\n输入报名商品的商品ID','SKU维度报名时必填，输入报名商品的SKUID','必填价格'])
    ws.append(['123456','222222',100]);b=BytesIO();wb.save(b)
    assert guard.validate_workbook(ready,b.getvalue())['ok']


def test_stale_export_does_not_override_current_erp_mapping(ready):
    from app.models.taobao_listing import TaobaoListing
    ready.add(TaobaoListing(taobao_item_id='123456',taobao_sku_id='222222',sku_code='P11'))
    ready.commit()
    assert guard.validate_workbook(ready,workbook([['123456','222222',20]]))['ok']


def test_official_merged_item_id_only_resolves_real_merge(ready):
    wb=openpyxl.Workbook();ws=wb.active;ws.append(['商品ID','SKUID','活动价'])
    ws.append(['123456','222222',100]);ws.append([None,'222222',100]);ws.merge_cells('A2:A3')
    b=BytesIO();wb.save(b)
    assert guard.validate_workbook(ready,b.getvalue())['checked_count']==2
    ws.unmerge_cells('A2:A3');b=BytesIO();wb.save(b)
    assert not guard.validate_workbook(ready,b.getvalue())['ok']


def test_dynamic_offsale_never_becomes_permanent_platform_exclusion(ready):
    from app.services.delisted_sku_service import add_delisted, get_delisted
    ready.execute(update(Product).values(listing_status='下架'));ready.commit()
    assert '222222' in get_delisted(ready)
    add_delisted(ready,['999999'])
    ready.execute(update(Product).values(listing_status='在售'));ready.commit()
    assert get_delisted(ready)=={'111111','333333','999999'}


def test_all_known_active_product_labels_remain_eligible(ready):
    from app.services.campaign_service import _erp_listed_product_codes
    for label in guard.ACTIVE:
        ready.execute(update(Product).values(listing_status=label));ready.commit()
        assert 'P' in _erp_listed_product_codes(ready)
        assert guard.validate_workbook(ready,workbook([['123456','222222',20]]))['ok']
